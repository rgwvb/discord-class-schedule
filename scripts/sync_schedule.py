import io, json, os, re, pathlib
from datetime import datetime
from zoneinfo import ZoneInfo
import requests
from openpyxl import load_workbook

URL=os.environ['ONEDRIVE_XLSX_URL']
OUT=pathlib.Path('data/schedule.json')
PEOPLE=['鍾曜遠','王奕文','吳承恩','林鶴翔','胡育祥','康思賢','郭致嘉','翁和緯']
DAYS={2:'1',3:'2',4:'3',5:'4',6:'5',7:'6'}
ROW_STARTS=[3,7,11,15,19,23,27,31]

def clean(v):
    if v is None:return ''
    return str(v).strip()

def maybe_num(v):
    if v in ('',None):return None
    try:
        n=float(v); return int(n) if n.is_integer() else n
    except:return None

def room_credit(a,b):
    sa,sb=clean(a),clean(b)
    na,nb=maybe_num(a),maybe_num(b)
    def looks_room(s):
        return bool(re.search(r'[A-Za-z?/]|\d{4,}',s))
    if looks_room(sa) and not looks_room(sb): return sa,nb
    if looks_room(sb) and not looks_room(sa): return sb,na
    if na is not None and nb is not None:
        if na>20 and nb<=10:return sa,nb
        if nb>20 and na<=10:return sb,na
    return sa,nb

def main():
    r=requests.get(URL,timeout=45); r.raise_for_status()
    wb=load_workbook(io.BytesIO(r.content),data_only=False)
    people={}
    for name in PEOPLE:
        if name not in wb.sheetnames: continue
        ws=wb[name]; days={str(i):[] for i in range(1,7)}
        for row in ROW_STARTS:
            for col,day in DAYS.items():
                title=clean(ws.cell(row,col).value)
                if not title or title in {'*','#REF!'}: continue
                teacher=clean(ws.cell(row+1,col).value)
                room,credits=room_credit(ws.cell(row+2,col).value,ws.cell(row+3,col).value)
                time=''
                raw=clean(ws.cell(row,1).value)
                if re.search(r'\d{1,2}:\d{2}',raw): time=raw
                if not time:
                    starts={3:'09:10–12:00',7:'09:10–12:00',11:'09:10–12:00',15:'13:10–16:00',19:'13:10–15:00',23:'15:10–17:00',27:'15:10–17:00',31:'13:10–16:00'}
                    time=starts.get(row,'')
                days[day].append({'time':time,'title':title,'teacher':teacher,'room':room,'credits':credits})
        people[name]=days
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps({'updated_at':datetime.now(ZoneInfo('Asia/Taipei')).isoformat(),'people':people},ensure_ascii=False,separators=(',',':')),encoding='utf-8')

if __name__=='__main__':main()
